"""Convert a depth-chart workbook export into the repo's capture format.

The depth-chart source can arrive two ways: scraped from the web page directly,
or exported to a spreadsheet. The spreadsheet form nests a stat sub-row under
every tier row and closes each team with a `<team> Leaders` block, and several
teams' final tier has no stat sub-row at all when none of its players recorded
stats -- so the Leaders title is the only reliable section boundary. Guessing
from blank rows drops players.

Emits the same `TEAM=<abbrev> ROWS=<n>` / `tier|PG|SG|SF|PF|C` format that
`depth_charts.py` parses, so the two capture routes converge on one file.

    python scripts/parse_depth_chart_workbook.py <workbook.xlsx> [>> capture.txt]
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

TEAMS = {
    'Atlanta Hawks': 'ATL', 'Boston Celtics': 'BOS', 'Brooklyn Nets': 'BRK',
    'Charlotte Hornets': 'CHO', 'Chicago Bulls': 'CHI', 'Cleveland Cavaliers': 'CLE',
    'Dallas Mavericks': 'DAL', 'Denver Nuggets': 'DEN', 'Detroit Pistons': 'DET',
    'Golden State Warriors': 'GSW', 'Houston Rockets': 'HOU', 'Indiana Pacers': 'IND',
    'Los Angeles Clippers': 'LAC', 'Los Angeles Lakers': 'LAL', 'Memphis Grizzlies': 'MEM',
    'Miami Heat': 'MIA', 'Milwaukee Bucks': 'MIL', 'Minnesota Timberwolves': 'MIN',
    'New Orleans Pelicans': 'NOP', 'New York Knicks': 'NYK', 'Oklahoma City Thunder': 'OKC',
    'Orlando Magic': 'ORL', 'Philadelphia Sixers': 'PHI', 'Philadelphia 76ers': 'PHI',
    'Phoenix Suns': 'PHO', 'Portland Trail Blazers': 'POR', 'Sacramento Kings': 'SAC',
    'San Antonio Spurs': 'SAS', 'Toronto Raptors': 'TOR', 'Utah Jazz': 'UTA',
    'Washington Wizards': 'WAS',
}
TIERS = {'Starters', 'Rotation', 'Lim PT'}


def parse(path: Path) -> list[tuple[str, str, list[str]]]:
    ws = load_workbook(path, data_only=True).worksheets[0]
    rows = [[('' if c is None else str(c).strip()) for c in r]
            for r in ws.iter_rows(values_only=True)]

    sections, cur = [], None
    for i, r in enumerate(rows):
        a = r[0]
        if a.endswith('Depth Chart'):
            name = a.rsplit('Depth Chart', 1)[0]
            name = name.split(' ', 1)[1].strip() if name[:1].isdigit() else name.strip()
            cur = {'team': TEAMS.get(name), 'name': name, 'start': i, 'end': None}
            sections.append(cur)
        elif a.endswith('Leaders') and cur and cur['end'] is None:
            cur['end'] = i

    out = []
    for s in sections:
        if s['team'] is None:
            raise SystemExit(f"unmapped team name: {s['name']!r} -- add it to TEAMS")
        if s['end'] is None:
            raise SystemExit(f"no Leaders block closing section for {s['name']!r}")
        body = [f"{r[0]}|" + "|".join(c if c else 'EMPTY' for c in r[1:6])
                for r in rows[s['start']:s['end']] if r[0] in TIERS]
        if not body:
            raise SystemExit(f"no tier rows found for {s['name']!r}")
        out.append((s['team'], s['name'], body))
    return out


def main() -> None:
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    parsed = parse(Path(sys.argv[1]))
    for team, name, body in parsed:
        n = sum(1 for b in body for c in b.split('|')[1:] if c != 'EMPTY')
        print(f"# {team}: {len(body)} rows, {n} players ({name})", file=sys.stderr)
        print(f"TEAM={team} ROWS={len(body)}")
        print("\n".join(body))
        print()
    total = sum(1 for _, _, b in parsed for ln in b
                for c in ln.split('|')[1:] if c != 'EMPTY')
    print(f"# parsed {len(parsed)} teams, {total} players", file=sys.stderr)


if __name__ == "__main__":
    main()
